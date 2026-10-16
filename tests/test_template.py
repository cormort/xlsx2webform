"""Fixed-column / variable-row collection templates (modelled on 表8 購建固定資產執行情形)."""
import io
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import openpyxl
import pytest
from fastapi.testclient import TestClient
import backend.main as main
from backend.main import app, sessions
from backend.xlsx_parser import process_xlsx, translate_formula

main.ADMIN_PASSWORD = "testpw"
c = TestClient(app)


def _template_bytes(fill_row=None):
    wb = openpyxl.Workbook()
    hidden = wb.active
    hidden.title = "舊版"
    hidden.sheet_state = "hidden"
    ws = wb.create_sheet("114非營業")
    ws.merge_cells("A1:F1"); ws["A1"] = "非營業特種基金114年度執行情形"
    ws["F2"] = "單位：元"
    ws.merge_cells("A3:A4"); ws["A3"] = "基金名稱"
    ws.merge_cells("B3:D3"); ws["B3"] = "可用預算數"
    ws["B4"], ws["C4"], ws["D4"] = "以前年度保留數", "本年度預算數", "合計"
    ws.merge_cells("E3:E4"); ws["E3"] = "決算數"
    ws.merge_cells("F3:F4"); ws["F3"] = "執行率"
    ws["G3"] = "檢查"
    ws["A5"] = "合計"
    for col in "BCDE":
        ws[f"{col}5"] = f"=SUM({col}6:{col}7)"
    ws["F5"] = "=E5/D5"
    for r, name in ((6, "甲基金"), (7, "乙基金")):
        ws[f"A{r}"] = name
        ws[f"D{r}"] = f"=SUM(B{r}:C{r})"
        ws[f"F{r}"] = f"=E{r}/D{r}"
        ws[f"G{r}"] = f'=IF(F{r}<0.8,"請記得查填","OK")'
    if fill_row:
        ws["B6"], ws["C6"], ws["E6"] = fill_row
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _isolate(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "DATA_DIR", tmp_path)
    monkeypatch.setattr(main, "PERSIST_FILES", {k: tmp_path / v.name for k, v in main.PERSIST_FILES.items()})
    monkeypatch.setattr(main, "UPLOAD_DIR", tmp_path)
    app.state.limiter.enabled = False
    yield
    for sid in list(sessions):
        if sessions[sid].name.startswith("tpl"):
            sessions.pop(sid); main.response_store.pop(sid, None)


def test_translate_formula():
    assert translate_formula("=SUM(B8:D8)", 4) == "=(B4+C4+D4)"
    assert translate_formula("=F8/E8", 4) == "=F4/E4"
    assert translate_formula("=SUM(B6:B7)-B3", 4) is None          # reaches into the header
    assert translate_formula('=IF(G8<0.8,"x","OK")', 4) is None
    assert translate_formula("=[1]大基金!$G$9", 4) is None


def test_parse_template(tmp_path):
    p = tmp_path / "t.xlsx"; p.write_bytes(_template_bytes())
    r = process_xlsx(str(p))
    m, j = r["metadata"], r["json"]
    assert m["sheet_name"] == "114非營業"                              # hidden first sheet skipped
    assert m["header_rows"] == [3, 4]
    assert m["caption"] == ["非營業特種基金114年度執行情形", "單位：元"]
    assert [x["value"] for x in j[0]] == ["基金名稱", "可用預算數／以前年度保留數", "可用預算數／本年度預算數",
                                          "可用預算數／合計", "決算數", "執行率", "檢查"]
    assert j[1][0]["value"] == "合計" and j[1][1]["_formula"] == "=(B2+B3)"
    assert j[2][3]["_formula"] == "=(B2+C2)" and j[2][5]["_formula"] == "=E2/D2"
    assert j[2][6].get("_locked")                                       # IF() not translatable
    assert [c["row"] for c in j[2][:1]] == [3]                          # renumbered for export
    # legacy mode keeps the whole sheet
    assert process_xlsx(str(p), header_row=0)["json"][0][0]["value"] == "非營業特種基金114年度執行情形"


def _admin():
    return c.post("/api/admin/login", json={"password": "testpw"}).json()["token"]


def test_template_flow_locks_import_consolidate_reparse():
    H = {"X-Admin-Token": _admin()}
    up = c.post("/api/upload-xlsx", files={"file": ("t.xlsx", _template_bytes())},
                data={"name": "tpl", "password": "pw"}, headers=H).json()
    sid = up["session_id"]
    assert up["metadata"]["header_rows"] == [3, 4]
    token = c.post(f"/api/sessions/{sid}/publish", headers=H).json()["share_token"]

    # respondent imports a filled copy of the same template → lines up with the combined header
    imp = c.post(f"/api/fill/{token}/parse-xlsx", files={"file": ("f.xlsx", _template_bytes((100, 200, 250)))}).json()
    assert imp["json"][0][1]["value"] == "可用預算數／以前年度保留數"
    assert [x["value"] for x in imp["json"][2][:3]] == ["甲基金", "100", "200"]

    # submit with a tampered label + header: locked cells are restored server-side
    data = imp["json"]
    data[0][0]["value"] = "HACK"
    data[2][0]["value"] = "改名基金"
    data.append([{"value": "自行新增基金"}, {"value": "5"}])
    assert c.post(f"/api/fill/{token}/submit", json={"data": data, "email": "f@x.com"}).status_code == 200
    stored = main.response_store[sid][0]["data"]
    assert stored[0][0]["value"] == "基金名稱" and stored[2][0]["value"] == "甲基金"
    assert stored[2][1]["value"] == "100"                               # respondent input kept

    # consolidation keeps only rows the respondent filled (甲基金 + the added row), not 合計/乙基金
    cons = c.get("/api/consolidate", headers=H).json()
    raws = sorted(row["fund_raw"] for g in cons["groups"] for a in g["agencies"] for f in a["funds"] for row in f["rows"])
    assert raws == ["甲基金", "自行新增基金"]

    # structure is frozen once responses exist
    assert c.post(f"/api/sessions/{sid}/reparse", json={"header_row": 1}, headers=H).status_code == 409


def test_reparse_changes_sheet_and_header():
    H = {"X-Admin-Token": _admin()}
    sid = c.post("/api/upload-xlsx", files={"file": ("t.xlsx", _template_bytes())},
                 data={"name": "tpl2"}, headers=H).json()["session_id"]
    r = c.post(f"/api/sessions/{sid}/reparse", json={"sheet_index": 1, "header_row": 1}, headers=H).json()
    assert r["metadata"]["header_rows"] == [1, 1] and r["json"][0][0]["value"] == "非營業特種基金114年度執行情形"
    assert c.post(f"/api/sessions/{sid}/reparse", json={"header_row": 0}, headers=H).status_code == 400
